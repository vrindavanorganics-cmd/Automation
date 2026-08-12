"""Word (.docx) operations via python-docx — real document manipulation."""
from __future__ import annotations

from pathlib import Path

from docx import Document


def create_document(path: str | Path, paragraphs: list[str], title: str | None = None) -> Path:
    doc = Document()
    if title:
        doc.add_heading(title, level=1)
    for para in paragraphs:
        doc.add_paragraph(para)
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(p))
    return p


def read_document(path: str | Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs)
