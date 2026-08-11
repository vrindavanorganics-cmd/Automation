"""Excel (.xlsx) operations via openpyxl — real spreadsheet manipulation,
not GUI automation.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


def create_workbook(path: str | Path, rows: list[list[Any]], sheet_name: str = "Sheet1") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(p)
    return p


def create_workbook_from_records(path: str | Path, records: list[dict], sheet_name: str = "Sheet1") -> Path:
    if not records:
        return create_workbook(path, [], sheet_name)
    headers = list(records[0].keys())
    rows = [headers] + [[rec.get(h, "") for h in headers] for rec in records]
    return create_workbook(path, rows, sheet_name)


def read_workbook(path: str | Path, sheet_name: str | None = None) -> list[list[Any]]:
    wb = load_workbook(str(path), data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    return [list(row) for row in ws.iter_rows(values_only=True)]


def read_workbook_as_records(path: str | Path, sheet_name: str | None = None) -> list[dict]:
    rows = read_workbook(path, sheet_name)
    if not rows:
        return []
    headers, *data_rows = rows
    return [dict(zip(headers, row)) for row in data_rows]


def append_row(path: str | Path, row: list[Any], sheet_name: str | None = None) -> None:
    wb = load_workbook(str(path))
    ws = wb[sheet_name] if sheet_name else wb.active
    ws.append(row)
    wb.save(str(path))


def update_cell(path: str | Path, cell: str, value: Any, sheet_name: str | None = None) -> None:
    wb = load_workbook(str(path))
    ws = wb[sheet_name] if sheet_name else wb.active
    ws[cell] = value
    wb.save(str(path))


def list_sheets(path: str | Path) -> list[str]:
    wb = load_workbook(str(path), read_only=True)
    return wb.sheetnames
